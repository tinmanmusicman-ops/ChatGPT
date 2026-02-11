import subprocess
import sys
import threading
import time
import os
import json
import re
import smtplib
import textwrap
from html import escape
from datetime import datetime, timedelta
from flask import Flask, send_from_directory, send_file, jsonify, request, redirect, abort
from werkzeug.exceptions import HTTPException
import requests
from pathlib import Path
from email.message import EmailMessage
from email.utils import formatdate
from typing import Optional
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, HRFlowable

# Import stub modules
from modules import (
    job_pipeline,
    fax_engine,
    thermostat,
    security_notify,
    config_manager,
    logger,
    task_scheduler,
)

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent
UI_DIR = BASE_DIR / "ui"
IMAGES_DIR = BASE_DIR.parent / "Images"
FIT_SITE_DIR = BASE_DIR.parent / "AI-Fit-Site"
DEFAULT_RESUME_OUTPUT_DIR = Path(r"C:\!!!!!!!!!!!!!!!!!!!!!!!!!Stuff")
BASE_OUTPUT_DIR = Path(os.environ.get("HSST_RESUME_OUTPUT_DIR", str(DEFAULT_RESUME_OUTPUT_DIR)))
BASE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
AI_BOTS_ROOT = BASE_DIR.parent / "ai-bots"
DASHBOARD_SCRIPT_PATH = AI_BOTS_ROOT / "Thermostats" / "scripts" / "Dashboard.py"
THERMOSTATS_WEB_DIR = AI_BOTS_ROOT / "Thermostats" / "Web"
GLOBAL_CONFIG_PATH = Path(
    os.environ.get("HSST_GLOBAL_CONFIG", str(AI_BOTS_ROOT / "shared" / "Global.json"))
)

DEFAULT_CORS_ORIGINS = {
    "http://localhost",
    "http://127.0.0.1",
    "http://localhost:80",
    "http://127.0.0.1:80",
}


def _allowed_origins():
    allowed = set(DEFAULT_CORS_ORIGINS)
    extra = os.environ.get("HSST_TOWER_ALLOWED_ORIGINS", "")
    for item in extra.split(","):
        item = item.strip()
        if item:
            allowed.add(item)
    return allowed


def _with_cors(resp):
    origin = request.headers.get("Origin", "").strip()
    if not origin:
        return resp

    allowed = _allowed_origins()
    if origin in allowed:
        resp.headers["Access-Control-Allow-Origin"] = origin
        resp.headers["Vary"] = "Origin"
        resp.headers["Access-Control-Allow-Methods"] = "POST, OPTIONS"
        resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return resp


def _append_cores_log(message: str):
    CORES_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().isoformat(timespec="seconds")
    entry = f"{timestamp} {message}\n"
    try:
        with open(CORES_LOG_PATH, "a", encoding="utf-8") as fh:
            fh.write(entry)
    except Exception:
        pass

CORES_DIR = BASE_DIR.parent / "CORES"
CORES_LOG_PATH = CORES_DIR / "cores.log"
CORES_EMPTY_TEMPLATE_PATH = CORES_DIR / "CORES_EMPTY_TEMPLATE.md"
CORES_COMPANY_TEMPLATE_PATH = CORES_DIR / "Company.md"
CORE_ID_FILE_PATTERN = re.compile(r"^CORE-US-(\d{4})-(\d{6})\.md$", flags=re.IGNORECASE)
CORES_REQUIRED_PROJECT_LABELS = [
    "Situation",
    "What I did",
    "Tools / systems",
    "Result",
    "Evidence / artifacts",
    "Notes / caveats",
]
PROJECT_PARSE_PROMPT = """Convert this real-world work story into a CORES project JSON object.

Return ONLY valid JSON with this exact shape:
{
  "projectTitle": "string (can be empty)",
  "sections": {
    "Situation": "string",
    "What I did": "string",
    "Tools / systems": "string",
    "Result": "string",
    "Evidence / artifacts": "string",
    "Notes / caveats": "string"
  }
}

Rules:
- No markdown.
- No commentary.
- Do not omit sections.
- Keep section names exactly as shown.
"""

COMPANY_PARSE_PROMPT = """Convert this real-world work story into a CORES company JSON object.

Return ONLY valid JSON with this exact shape:
{
  "companyName": "string",
  "role": "string",
  "signals": "string",
  "firstProject": {
    "projectTitle": "string (can be empty)",
    "sections": {
      "Situation": "string",
      "What I did": "string",
      "Tools / systems": "string",
      "Result": "string",
      "Evidence / artifacts": "string",
      "Notes / caveats": "string"
    }
  }
}

Rules:
- No markdown.
- No commentary.
- Do not omit fields.
- Keep section names exactly as shown.
- "companyName", "role", and "signals" must not be empty.
- "firstProject.sections" must include all required fields and each field must be non-empty.
"""


def _load_company_template_markdown() -> str:
    if not CORES_COMPANY_TEMPLATE_PATH.exists():
        raise ValueError(f"Company template not found: {CORES_COMPANY_TEMPLATE_PATH.name}")
    template_markdown = (
        CORES_COMPANY_TEMPLATE_PATH.read_text(encoding="utf-8")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )
    if not template_markdown:
        raise ValueError(f"Company template is empty: {CORES_COMPANY_TEMPLATE_PATH.name}")
    return template_markdown


def _build_company_formatter_block(template_markdown: str) -> str:
    return (
        "Everything above this line is the company story.\n\n"
        "Return markdown only.\n"
        "Do not add commentary.\n"
        "Do not add explanations.\n"
        "Do not add intro text.\n"
        "Do not add headings beyond this structure.\n"
        "Do not convert bullet markers to headings.\n"
        "Return the final company block inside ONE fenced code block using triple backticks.\n"
        "No text before the fenced block.\n"
        "No text after the fenced block.\n"
        "Keep literal markdown characters exactly as shown.\n"
        "Use '-' dash bullets exactly; do not use Unicode bullets.\n"
        "Do NOT create additional bullet items inside section content.\n"
        "After each required label, write plain text only (no nested '-' bullets, no numbered lists).\n"
        "For section content lines, never start with '-', '*', '#', or a numbered list marker like '1.'.\n"
        "If source has multiple points, combine them into one plain sentence separated by semicolons.\n"
        "Do not add extra bullet lines under any section.\n"
        "The first heading line must start with exactly: ## COMPANY:\n\n"
        "Use this exact company framework with all content populated:\n\n"
        f"{template_markdown}\n"
    )

CORES_FORMATTER_BLOCK = """Everything above this line is the project story.

Return markdown only.
Do not add commentary.
Do not add explanations.
Do not add intro text.
Do not add headings beyond this structure.
Do not convert bullet markers to headings.
Return the final project inside ONE fenced code block using triple backticks.
No text before the fenced block.
No text after the fenced block.
Keep literal markdown characters exactly as shown.
Use '-' dash bullets exactly; do not use Unicode bullets.
Do NOT create additional bullet items inside section content.
After each required label, write plain text only (no nested '-' bullets, no numbered lists).
For section content lines, never start with '-', '*', '#', or a numbered list marker like '1.'.
If source has multiple points, combine them into one plain sentence separated by semicolons.
Do not add extra bullet lines under any section.

Project title must be blank.
The first line must be exactly:
### PROJECT:

Use this exact structure with all content populated:

### PROJECT:

- Situation:

- What I did:

- Tools / systems:

- Result:

- Evidence / artifacts:

- Notes / caveats:
"""

ASCII_CHAR_REPLACEMENTS = {
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u2013": "-",
    "\u2014": "-",
    "\u2022": "-",
    "\u2026": "...",
    "\u00a0": " ",
}

RESUME_IMPORT_PROMPT = """TASK:
Parse the provided PDF or DOCX resume and populate the provided EMPTY CORES-MD TEMPLATE verbatim as a starting point. Every heading, label, and section order in the template is locked; do not invent or reorder anything.

ROLE:
You are a deterministic parser and normalizer, not a resume writer.
Do NOT embellish, optimize, or invent content.
Do NOT introduce new sections.
Do NOT remove required sections.
Preserve meaning; normalize structure.

INPUTS:
1) A resume source text extracted from PDF/DOCX.
2) An EMPTY CORES-MD template (authoritative schema).

OUTPUT:
A single, fully populated CORES-MD file that mirrors the template structure exactly. Use the template as the sole reference for headings, labels, and ordering—do not guess or rely on any other schema.

HARD INVARIANTS (FAIL FAST):
-- Use ONLY sections present in the empty template exactly as written.
-- DO NOT add non-standard sections.
-- ALL items under '# SKILLS' MUST begin with '- '.
-- If content does not clearly map, leave the field blank rather than guessing.
-- No placeholders like 'XXX', 'TBD', or commentary.
-- No em dashes.
-- Use ASCII characters only.
-- No lines longer than 120 characters.
-- Headings must remain exactly as provided in the template and in the same order.
-- On any heading mismatch, stop immediately and respond with:
   SCHEMA_VIOLATION: HEADINGS_MISMATCH
-- Bullets must be one item per line.

PARSING RULES:
- HUMAN:
  - Name, Location, Availability: extract verbatim if present.
  - If Email/Phone/PURL not present in source, leave blank.
- SUMMARY:
  - Merge professional summary, working style, and strengths into a concise factual summary.
  - No marketing language.
- SKILLS:
  - Extract explicit skills, tools, or strengths.
  - Normalize into short noun phrases.
  - One skill per bullet.
- EXPERIENCE:
  - Each COMPANY block represents one employer.
  - ROLE, SIGNALS, PROJECT blocks must originate directly from the source content.
  - PROJECT descriptions should reflect actual work described, not inferred impact.
- ACHIEVEMENTS:
  - Include only concrete outcomes or transitions described in the source.

ERROR HANDLING:
- If a schema violation would occur, STOP and output exactly:
SCHEMA_VIOLATION: <short reason>

FINAL OUTPUT:
- Output ONLY the completed CORES-MD content.
- No explanations.
- No commentary.
- No markdown fences.
"""


def _resolve_core_markdown_path(core_id: str) -> Path:
    clean_core_id = str(core_id or "").strip()
    if not clean_core_id:
        raise ValueError("coreId is required")
    if clean_core_id in {".", ".."}:
        raise ValueError("Invalid coreId format")
    if re.search(r"[\\/:*?\"<>|]", clean_core_id):
        raise ValueError("Invalid coreId format")

    base = CORES_DIR.resolve()
    path = (base / f"{clean_core_id}.md").resolve()
    if not str(path).startswith(str(base)):
        raise PermissionError("Path traversal blocked")
    return path


def _next_core_id_for_year(target_year: int) -> str:
    max_sequence = 0
    if CORES_DIR.exists():
        for path in CORES_DIR.iterdir():
            if not path.is_file():
                continue
            match = CORE_ID_FILE_PATTERN.match(path.name)
            if not match:
                continue
            year_value = int(match.group(1))
            seq_value = int(match.group(2))
            if year_value == target_year and seq_value > max_sequence:
                max_sequence = seq_value
    next_sequence = max_sequence + 1
    return f"CORE-US-{target_year}-{next_sequence:06d}"


def _with_core_id_in_template(template_markdown: str, core_id: str) -> str:
    template = str(template_markdown or "").replace("\r\n", "\n").replace("\r", "\n")
    updated = re.sub(r"(?im)^ID:\s*.*$", f"ID: {core_id}", template, count=1)
    if updated == template:
        raise ValueError("Template is missing required 'ID:' line.")
    return updated.strip() + "\n"


def _parse_project_block_sections(project_block: str) -> dict:
    normalized = str(project_block or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise ValueError("Project block is empty")

    lines = normalized.split("\n")
    header_match = re.match(r"^### PROJECT:\s*(.*)$", lines[0].strip())
    if not header_match:
        raise ValueError("Project block must start with '### PROJECT:'")

    project_title = header_match.group(1).strip()
    sections = {label: [] for label in CORES_REQUIRED_PROJECT_LABELS}
    current_label = None
    next_required_index = 0

    for line in lines[1:]:
        marker_match = re.match(r"^- ([^:]+):\s*(.*)$", line.strip())
        if marker_match:
            label = marker_match.group(1).strip()
            initial_value = marker_match.group(2)

            if label not in CORES_REQUIRED_PROJECT_LABELS:
                raise ValueError(f"Unsupported section label: {label}")
            if next_required_index >= len(CORES_REQUIRED_PROJECT_LABELS):
                raise ValueError("Unexpected additional section marker")
            expected_label = CORES_REQUIRED_PROJECT_LABELS[next_required_index]
            if label != expected_label:
                raise ValueError(f"Section out of order: expected '{expected_label}', got '{label}'")

            current_label = label
            sections[current_label].append(initial_value)
            next_required_index += 1
            continue

        if current_label is None:
            if not line.strip():
                continue
            raise ValueError("Unexpected content before first section marker")
        sections[current_label].append(line)

    if next_required_index != len(CORES_REQUIRED_PROJECT_LABELS):
        missing = CORES_REQUIRED_PROJECT_LABELS[next_required_index:]
        raise ValueError(f"Missing required sections: {', '.join(missing)}")

    cleaned_sections = {
        label: "\n".join(value_lines).strip()
        for label, value_lines in sections.items()
    }
    return {
        "projectTitle": project_title,
        "sections": cleaned_sections,
    }


def _build_project_block(project_title: str, sections: dict) -> str:
    title = str(project_title or "").strip()
    lines = [f"### PROJECT: {title}" if title else "### PROJECT:", ""]
    for label in CORES_REQUIRED_PROJECT_LABELS:
        value = str(sections.get(label, "") or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        lines.append(f"- {label}:")
        if value:
            lines.extend(value.split("\n"))
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def _extract_first_json_object(text: str) -> str:
    raw = str(text or "").strip()
    if not raw:
        raise ValueError("AI returned empty content")

    fenced = re.search(r"```(?:json)?\s*(\{[\s\S]*?\})\s*```", raw, flags=re.IGNORECASE)
    if fenced:
        return fenced.group(1).strip()

    start = raw.find("{")
    if start < 0:
        raise ValueError("No JSON object found in AI response")

    depth = 0
    in_string = False
    escape_next = False
    for idx in range(start, len(raw)):
        ch = raw[idx]
        if escape_next:
            escape_next = False
            continue
        if ch == "\\":
            escape_next = True
            continue
        if ch == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return raw[start : idx + 1]

    raise ValueError("Unclosed JSON object in AI response")


def _normalize_project_json(project_json: dict) -> dict:
    if not isinstance(project_json, dict):
        raise ValueError("AI JSON must be an object")

    project_title = str(project_json.get("projectTitle") or project_json.get("title") or "").strip()
    sections_raw = project_json.get("sections")
    if not isinstance(sections_raw, dict):
        raise ValueError("AI JSON missing 'sections' object")

    alias_map = {
        "situation": "Situation",
        "what i did": "What I did",
        "tools / systems": "Tools / systems",
        "tools/systems": "Tools / systems",
        "result": "Result",
        "evidence / artifacts": "Evidence / artifacts",
        "evidence/artifacts": "Evidence / artifacts",
        "notes / caveats": "Notes / caveats",
        "notes/caveats": "Notes / caveats",
    }

    normalized_sections = {}
    for raw_key, raw_value in sections_raw.items():
        key = str(raw_key or "").strip().lower()
        key = re.sub(r"\s*\/\s*", "/", key)
        key = re.sub(r"\s+", " ", key)
        canonical = alias_map.get(key)
        if not canonical:
            continue
        value = str(raw_value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
        if not value:
            raise ValueError(f"AI section '{canonical}' is empty")
        normalized_sections[canonical] = value

    missing = [label for label in CORES_REQUIRED_PROJECT_LABELS if label not in normalized_sections]
    if missing:
        raise ValueError(f"AI JSON missing required sections: {', '.join(missing)}")

    return {
        "projectTitle": project_title,
        "sections": normalized_sections,
    }


def _normalize_company_json(company_json: dict) -> dict:
    if not isinstance(company_json, dict):
        raise ValueError("AI JSON must be an object")

    company_name = str(company_json.get("companyName") or company_json.get("name") or "").strip()
    role = str(company_json.get("role") or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    signals = str(company_json.get("signals") or "").replace("\r\n", "\n").replace("\r", "\n").strip()

    if not company_name:
        raise ValueError("AI JSON missing companyName")
    if not role:
        raise ValueError("AI JSON missing role")
    if not signals:
        raise ValueError("AI JSON missing signals")

    first_project_raw = company_json.get("firstProject")
    if not isinstance(first_project_raw, dict):
        first_project_raw = company_json.get("project")
    if not isinstance(first_project_raw, dict):
        raise ValueError("AI JSON missing firstProject object")

    first_project = _normalize_project_json(first_project_raw)
    return {
        "companyName": company_name,
        "role": role,
        "signals": signals,
        "firstProject": first_project,
    }


def _extract_pdf_text(file_bytes: bytes) -> str:
    PdfReader = None
    try:
        from pypdf import PdfReader as _PdfReader  # type: ignore
        PdfReader = _PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader as _PdfReader  # type: ignore
            PdfReader = _PdfReader
        except Exception as exc:
            raise ValueError("PDF parsing dependency is missing (install pypdf or PyPDF2).") from exc

    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Failed to parse PDF: {exc}") from exc

    chunks = []
    for page in reader.pages:
        text = str(page.extract_text() or "").strip()
        if text:
            chunks.append(text)

    merged = "\n\n".join(chunks).strip()
    if not merged:
        raise ValueError("No extractable text found in PDF.")
    return merged


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except Exception as exc:
        raise ValueError("DOCX parsing dependency is missing (install python-docx).") from exc

    try:
        doc = DocxDocument(BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Failed to parse DOCX: {exc}") from exc

    chunks = []
    for paragraph in doc.paragraphs:
        text = str(paragraph.text or "").strip()
        if text:
            chunks.append(text)

    merged = "\n\n".join(chunks).strip()
    if not merged:
        raise ValueError("No extractable text found in DOCX.")
    return merged


def _extract_heading_lines(markdown_text: str) -> list[str]:
    normalized = str(markdown_text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized.split("\n")]
    return [line for line in lines if re.match(r"^#{1,6}\s+\S", line)]


def _normalize_heading_line(heading_line: str) -> str:
    match = re.match(r"^(#{1,6})\s+(.*)$", str(heading_line or "").strip())
    if not match:
        return str(heading_line or "").strip()
    hashes = match.group(1)
    text = re.sub(r"\s+", " ", match.group(2).strip())
    return f"{hashes} {text}"


def _is_dynamic_template_heading(heading_line: str) -> bool:
    heading = _normalize_heading_line(heading_line)
    upper = heading.upper()
    if "XXX" in upper:
        return True
    if re.match(r"^#{2,6}\s+COMPANY:\s*", upper):
        return True
    if re.match(r"^#{2,6}\s+PROJECT:\s*", upper):
        return True
    return False


def _normalize_common_ascii(text: str) -> str:
    normalized = str(text or "")
    for src, dst in ASCII_CHAR_REPLACEMENTS.items():
        normalized = normalized.replace(src, dst)
    return normalized


def _describe_non_ascii(text: str, limit: int = 8) -> str:
    found = []
    seen = set()
    for ch in str(text or ""):
        code = ord(ch)
        if code <= 127:
            continue
        key = (code, ch)
        if key in seen:
            continue
        seen.add(key)
        found.append(f"U+{code:04X} '{ch}'")
        if len(found) >= limit:
            break
    return ", ".join(found)


def _wrap_line(text: str, width: int) -> list[str]:
    if len(text) <= width:
        return [text]
    wrapped = textwrap.wrap(
        text,
        width=width,
        break_long_words=False,
        break_on_hyphens=False,
        replace_whitespace=False,
        drop_whitespace=True,
    )
    return wrapped or [text]


def _enforce_resume_import_line_length(markdown_text: str, max_len: int = 120) -> str:
    normalized = str(markdown_text or "").replace("\r\n", "\n").replace("\r", "\n")
    out_lines = []
    in_skills = False

    for line in normalized.split("\n"):
        stripped_line = line.strip()
        if re.match(r"^#\s+SKILLS\s*$", stripped_line, flags=re.IGNORECASE):
            in_skills = True
        elif re.match(r"^#\s+\S+", stripped_line) and not re.match(
            r"^#\s+SKILLS\s*$", stripped_line, flags=re.IGNORECASE
        ):
            in_skills = False

        if len(line) <= max_len:
            out_lines.append(line)
            continue

        # Keep heading lines unchanged so heading text remains exactly as output.
        if stripped_line.startswith("#"):
            out_lines.append(line)
            continue

        leading_ws = re.match(r"^[ \t]*", line).group(0)
        content = line[len(leading_ws) :]
        if not content.strip():
            out_lines.append(line)
            continue

        if in_skills and content.startswith("- "):
            prefix = f"{leading_ws}- "
            skill_text = content[2:].strip()
            width = max_len - len(prefix)
            if width < 8:
                out_lines.append(line)
                continue
            for chunk in _wrap_line(skill_text, width):
                out_lines.append(f"{prefix}{chunk}")
            continue

        width = max_len - len(leading_ws)
        if width < 8:
            out_lines.append(line)
            continue
        for chunk in _wrap_line(content.strip(), width):
            out_lines.append(f"{leading_ws}{chunk}")

    return "\n".join(out_lines).strip() + "\n"


def _validate_resume_import_output(template_markdown: str, output_markdown: str) -> None:
    template = str(template_markdown or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    output = str(output_markdown or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not template:
        raise ValueError("Template markdown is empty.")
    if not output:
        raise ValueError("AI returned empty markdown.")
    if output.startswith("SCHEMA_VIOLATION:"):
        raise ValueError(output)
    if "```" in output:
        raise ValueError("Output must not include markdown code fences.")
    if "—" in output:
        raise ValueError("Output contains a forbidden em dash.")
    if re.search(r"\b(?:XXX|TBD)\b", output, flags=re.IGNORECASE):
        raise ValueError("Output contains forbidden placeholders.")
    if any(ord(ch) > 127 for ch in output):
        detail = _describe_non_ascii(output) or "unknown characters"
        raise ValueError(f"Output contains non-ASCII characters: {detail}.")

    lines = output.split("\n")
    for idx, line in enumerate(lines, start=1):
        if len(line) > 120:
            raise ValueError(f"Line {idx} exceeds 120 characters.")

    template_headings = _extract_heading_lines(template)
    output_headings = _extract_heading_lines(output)
    normalized_output = [_normalize_heading_line(line) for line in output_headings]
    required_template_headings = [
        _normalize_heading_line(line)
        for line in template_headings
        if not _is_dynamic_template_heading(line)
    ]

    output_index = 0
    for required_heading in required_template_headings:
        found_index = -1
        for idx in range(output_index, len(normalized_output)):
            if normalized_output[idx] == required_heading:
                found_index = idx
                break
        if found_index < 0:
            output_preview = ", ".join(normalized_output[:14])
            raise ValueError(
                "Heading mismatch with template. "
                f"Missing required heading: {required_heading}. "
                f"Output headings seen: {output_preview}"
            )
        output_index = found_index + 1

    skills_match = re.search(
        r"^#\s+SKILLS\s*$([\s\S]*?)(?=^#\s+\S|\Z)",
        output,
        flags=re.MULTILINE | re.IGNORECASE,
    )
    if skills_match:
        skills_block = skills_match.group(1)
        for line in skills_block.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            if not stripped.startswith("- "):
                raise ValueError("All # SKILLS items must start with '- '.")


def _validate_cores_markdown(markdown_text: str) -> None:
    normalized = str(markdown_text or "").replace("\r\n", "\n").replace("\r", "\n")
    if not normalized.strip():
        raise ValueError("Markdown cannot be empty")
    if not re.search(r"^#\s+CORES\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing # CORES header")
    if not re.search(r"^#\s+HUMAN\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing # HUMAN header")
    if not re.search(r"^#\s+SUMMARY\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing # SUMMARY header")
    if not re.search(r"^#\s+GLOBAL\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing # GLOBAL header")
    if not re.search(r"^#\s+EXPERIENCE\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing # EXPERIENCE header")
    if not re.search(r"^##\s+ACHIEVEMENTS\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE):
        raise ValueError("Missing ## ACHIEVEMENTS section")

    summary_match = re.search(r"^#\s+SUMMARY\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE)
    global_match = re.search(r"^#\s+GLOBAL\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE)
    certifications_match = re.search(
        r"^#\s+CERTIFICATIONS\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE
    )
    experience_match = re.search(r"^#\s+EXPERIENCE\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE)
    education_match = re.search(
        r"^#\s*EDUCATION\s*:?\s*$", normalized, flags=re.MULTILINE | re.IGNORECASE
    )
    if not (summary_match and global_match and experience_match):
        raise ValueError("Required section headers not found")
    if certifications_match and education_match:
        if not (
            summary_match.start()
            < global_match.start()
            < certifications_match.start()
            < experience_match.start()
            < education_match.start()
        ):
            raise ValueError(
                "Header order invalid: expected # SUMMARY -> # GLOBAL -> # Certifications -> # EXPERIENCE -> # Education"
            )
    elif certifications_match:
        if not (summary_match.start() < global_match.start() < certifications_match.start() < experience_match.start()):
            raise ValueError(
                "Header order invalid: expected # SUMMARY -> # GLOBAL -> # Certifications -> # EXPERIENCE"
            )
    elif education_match:
        if not (summary_match.start() < global_match.start() < experience_match.start() < education_match.start()):
            raise ValueError(
                "Header order invalid: expected # SUMMARY -> # GLOBAL -> # EXPERIENCE -> # Education"
            )
    elif not (summary_match.start() < global_match.start() < experience_match.start()):
        raise ValueError("Header order invalid: expected # SUMMARY -> # GLOBAL -> # EXPERIENCE")

    companies = list(re.finditer(r"^##\s+COMPANY:\s*(.+)$", normalized, flags=re.MULTILINE | re.IGNORECASE))
    if not companies:
        raise ValueError("No company blocks found")

    for index, company_match in enumerate(companies):
        block_start = company_match.start()
        block_end = companies[index + 1].start() if index + 1 < len(companies) else len(normalized)
        company_block = normalized[block_start:block_end]
        company_name = company_match.group(1).strip() or f"index {index}"

        if not re.search(r"^###\s+ROLE:\s*$", company_block, flags=re.MULTILINE | re.IGNORECASE):
            raise ValueError(f"Company '{company_name}' missing ### ROLE: header")
        if not re.search(r"^###\s+SIGNALS:\s*$", company_block, flags=re.MULTILINE | re.IGNORECASE):
            raise ValueError(f"Company '{company_name}' missing ### SIGNALS: header")

        projects = list(re.finditer(r"^###\s+PROJECT:\s*.*$", company_block, flags=re.MULTILINE | re.IGNORECASE))
        if not projects:
            raise ValueError(f"Company '{company_name}' has no ### PROJECT: blocks")

        for proj_idx, proj_match in enumerate(projects):
            proj_start = proj_match.start()
            proj_end = projects[proj_idx + 1].start() if proj_idx + 1 < len(projects) else len(company_block)
            project_block = company_block[proj_start:proj_end]

            for label in CORES_REQUIRED_PROJECT_LABELS:
                if not re.search(
                    rf"^-\s*{re.escape(label)}:\s*",
                    project_block,
                    flags=re.MULTILINE | re.IGNORECASE,
                ):
                    raise ValueError(
                        f"Company '{company_name}' project {proj_idx + 1} missing section '{label}'"
                    )


def _markdown_to_pdf_bytes(markdown: str) -> bytes:
    def inline_markup(text: str) -> str:
        safe = escape(text or "", quote=False)
        safe = re.sub(
            r"\[([^\]]+)\]\(([^)]+)\)",
            r'<u><font color="#2f6fd6">\1</font></u> <font color="#666666">(\2)</font>',
            safe,
        )
        safe = re.sub(r"`([^`]+)`", r'<font name="Courier">\1</font>', safe)
        safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
        safe = re.sub(r"__(.+?)__", r"<b>\1</b>", safe)
        safe = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", safe)
        safe = re.sub(r"(?<!_)_(?!\s)(.+?)(?<!\s)_(?!_)", r"<i>\1</i>", safe)
        return safe

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "MD_Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceBefore=2,
        spaceAfter=6,
    )
    heading_styles = {
        1: ParagraphStyle("MD_H1", parent=styles["Heading1"], fontSize=18, leading=22, spaceBefore=10, spaceAfter=8),
        2: ParagraphStyle("MD_H2", parent=styles["Heading2"], fontSize=16, leading=20, spaceBefore=9, spaceAfter=7),
        3: ParagraphStyle("MD_H3", parent=styles["Heading3"], fontSize=14, leading=18, spaceBefore=8, spaceAfter=6),
        4: ParagraphStyle("MD_H4", parent=styles["Heading4"], fontSize=12, leading=16, spaceBefore=7, spaceAfter=5),
        5: ParagraphStyle("MD_H5", parent=styles["Heading5"], fontSize=11, leading=15, spaceBefore=6, spaceAfter=4),
        6: ParagraphStyle("MD_H6", parent=styles["Heading6"], fontSize=10, leading=14, spaceBefore=5, spaceAfter=3),
    }
    quote_style = ParagraphStyle(
        "MD_Quote",
        parent=body_style,
        leftIndent=18,
        textColor=colors.HexColor("#444444"),
        italic=True,
    )
    list_style = ParagraphStyle(
        "MD_List",
        parent=body_style,
        leftIndent=18,
        firstLineIndent=0,
        spaceBefore=1,
        spaceAfter=2,
    )
    code_style = ParagraphStyle(
        "MD_Code",
        parent=styles["Code"],
        fontName="Courier",
        fontSize=8.5,
        leading=10.5,
        leftIndent=8,
        rightIndent=8,
        backColor=colors.whitesmoke,
        borderPadding=6,
        spaceBefore=4,
        spaceAfter=8,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title="CORE PDF",
    )

    story = []
    paragraph_lines = []
    list_items = []
    in_code_block = False
    code_lines = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines).strip()
        paragraph_lines.clear()
        if text:
            story.append(Paragraph(inline_markup(text), body_style))

    def flush_list() -> None:
        if not list_items:
            return
        ordered_counter = 0
        for list_type, indent_level, list_text in list_items:
            if list_type == "ol":
                ordered_counter += 1
                bullet = f"{ordered_counter}."
            else:
                bullet = "•"
            level_indent = max(indent_level, 0) * 12
            item_style = ParagraphStyle(
                "MD_List_Item",
                parent=list_style,
                leftIndent=list_style.leftIndent + level_indent,
            )
            story.append(Paragraph(inline_markup(list_text), item_style, bulletText=bullet))
        story.append(Spacer(1, 2))
        list_items.clear()

    def flush_code_block() -> None:
        if not code_lines:
            return
        code_text = "\n".join(code_lines)
        code_lines.clear()
        story.append(Preformatted(code_text, code_style))

    for raw_line in markdown.splitlines():
        stripped = raw_line.strip()
        fence_match = re.match(r"^```", stripped)
        if fence_match:
            if in_code_block:
                flush_code_block()
                in_code_block = False
            else:
                flush_paragraph()
                flush_list()
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw_line.rstrip("\n"))
            continue

        if not stripped:
            flush_paragraph()
            flush_list()
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            flush_list()
            level = min(len(heading_match.group(1)), 6)
            heading_text = heading_match.group(2).strip()
            story.append(Paragraph(inline_markup(heading_text), heading_styles[level]))
            continue

        if re.match(r"^([-*_])\1{2,}$", stripped):
            flush_paragraph()
            flush_list()
            story.append(HRFlowable(width="100%", color=colors.HexColor("#888888"), thickness=0.6))
            story.append(Spacer(1, 6))
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            flush_list()
            quote_text = stripped.lstrip(">").strip()
            if quote_text:
                story.append(Paragraph(inline_markup(quote_text), quote_style))
            continue

        unordered_match = re.match(r"^(\s*)[-+*]\s+(.+)$", raw_line)
        ordered_match = re.match(r"^(\s*)\d+\.\s+(.+)$", raw_line)
        if unordered_match or ordered_match:
            flush_paragraph()
            if unordered_match:
                list_type = "ul"
                leading_spaces, list_text = unordered_match.group(1), unordered_match.group(2)
            else:
                list_type = "ol"
                leading_spaces, list_text = ordered_match.group(1), ordered_match.group(2)
            indent_level = int(len(leading_spaces) / 2)
            if list_items and list_items[-1][0] != list_type:
                flush_list()
            list_items.append((list_type, indent_level, list_text.strip()))
            continue

        flush_list()
        paragraph_lines.append(stripped)

    if in_code_block:
        flush_code_block()
    flush_paragraph()
    flush_list()

    if not story:
        story.append(Paragraph("(empty document)", body_style))

    doc.build(story)
    return buffer.getvalue()


# -------------------------
# Serve the Control Tower UI
# -------------------------
@app.route("/", strict_slashes=False)
def root():
    _maybe_notify_fit_site_visit(page="/")
    return redirect("/tower", code=302)


@app.route("/tower", strict_slashes=False)
def tower_ui():
    _maybe_notify_fit_site_visit(page="/tower")
    return send_from_directory(UI_DIR, "index.html")


@app.route("/health")
def health():
    return {"status": "ok"}, 200


_global_config_cache = None
_global_config_mtime = None


def _load_global_config() -> dict:
    global _global_config_cache, _global_config_mtime

    try:
        stat = GLOBAL_CONFIG_PATH.stat()
    except FileNotFoundError:
        return {}

    if _global_config_cache is not None and _global_config_mtime == stat.st_mtime:
        return _global_config_cache or {}

    try:
        data = json.loads(GLOBAL_CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        data = {}

    _global_config_cache = data or {}
    _global_config_mtime = stat.st_mtime
    return _global_config_cache


def _send_gmail_message(*, to_address: str, subject: str, body: str, reply_to: Optional[str] = None) -> None:
    cfg = _load_global_config()
    gmail_user = (os.environ.get("GMAIL_USER") or cfg.get("gmail_user") or "").strip()
    gmail_app_password = (os.environ.get("GMAIL_APP_PASSWORD") or cfg.get("gmail_app_password") or "").strip()

    if not gmail_user or not gmail_app_password:
        raise RuntimeError("Missing Gmail credentials (GMAIL_USER/GMAIL_APP_PASSWORD or Global.json keys)")

    safe_subject = re.sub(r"[\r\n]+", " ", subject or "").strip()[:220] or "Contact request"

    msg = EmailMessage()
    msg["From"] = gmail_user
    msg["To"] = to_address
    msg["Subject"] = safe_subject
    msg["Date"] = formatdate(localtime=True)
    if reply_to:
        msg["Reply-To"] = reply_to
    msg.set_content(body or "")

    with smtplib.SMTP("smtp.gmail.com", 587, timeout=25) as smtp:
        smtp.ehlo()
        smtp.starttls()
        smtp.ehlo()
        smtp.login(gmail_user, gmail_app_password)
        smtp.send_message(msg)


def _looks_like_email(value: str) -> bool:
    s = (value or "").strip()
    if not s or len(s) > 254:
        return False
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", s))


_contact_last_sent_at = {}
_visit_last_notified_at = {}


# -------------------------
# Serve the AI-Fit-Site static page(s)
# -------------------------
@app.route("/AI-Fit-Site", strict_slashes=False)
def fit_site_index():
    _maybe_notify_fit_site_visit(page="Job.html")
    return send_from_directory(FIT_SITE_DIR, "Job.html")


@app.route("/manage", strict_slashes=False)
@app.route("/AI-Fit-Site/manage", strict_slashes=False)
@app.route("/AI-Fit-Site/manage.html", strict_slashes=False)
def fit_site_manage():
    return send_from_directory(FIT_SITE_DIR, "manage.html")


@app.route("/AI-Fit-Site/ai_context.pdf")
def fit_site_context_pdf():
    source_name = "CORE-US-2026-000001.md"
    requested_path = (CORES_DIR / source_name).resolve()
    base = CORES_DIR.resolve()

    try:
        if not str(requested_path).startswith(str(base)):
            abort(403)
        if not requested_path.exists():
            abort(404)

        markdown = requested_path.read_text(encoding="utf-8")
        pdf_bytes = _markdown_to_pdf_bytes(markdown)
        return send_file(
            BytesIO(pdf_bytes),
            mimetype="application/pdf",
            download_name="CORE-US-2026-000001.pdf",
        )
    except HTTPException:
        raise
    except Exception:
        return jsonify({"status": "error", "error": "PDF conversion failed"}), 500


@app.route("/AI-Fit-Site/<path:filename>")
def fit_site_static(filename):
    if str(filename or "").strip().lower() == "job.html":
        _maybe_notify_fit_site_visit(page="Job.html")
    return send_from_directory(FIT_SITE_DIR, filename)


def _is_obvious_bot_user_agent(user_agent: str) -> bool:
    ua = (user_agent or "").lower()
    if not ua:
        return False
    bot_markers = (
        "bot",
        "spider",
        "crawler",
        "slurp",
        "headless",
        "preview",
        "pingdom",
        "uptimerobot",
        "httpclient",
        "python-requests",
        "curl",
        "wget",
    )
    return any(marker in ua for marker in bot_markers)


def _maybe_notify_fit_site_visit(*, page: str) -> None:
    enabled = os.environ.get("HSST_VISIT_NOTIFY", "1").strip().lower() not in {"0", "false", "no"}
    if not enabled:
        return

    ip = (request.remote_addr or "").strip() or "unknown"
    ua = (request.headers.get("User-Agent", "") or "").strip()
    if _is_obvious_bot_user_agent(ua):
        return

    now = time.time()
    last = _visit_last_notified_at.get(ip)
    if last is not None and now - last < 6 * 60 * 60:
        return

    # Set before sending to avoid races on parallel requests.
    _visit_last_notified_at[ip] = now

    to_address = os.environ.get("HSST_VISIT_NOTIFY_TO", "timinman2024@gmail.com").strip()
    subject = f"AI-Fit-Site visit: {page}"
    referer = (request.headers.get("Referer", "") or "").strip()
    origin = (request.headers.get("Origin", "") or "").strip()
    path = (request.full_path or request.path or "").strip()

    body = (
        "AI-Fit-Site visit detected\n\n"
        f"Time: {datetime.now().isoformat(timespec='seconds')}\n"
        f"Page: {page}\n"
        f"Path: {path}\n"
        f"Remote: {ip}\n"
        f"Origin: {origin}\n"
        f"Referer: {referer}\n"
        f"User-Agent: {ua}\n"
    )

    if os.environ.get("HSST_VISIT_NOTIFY_DRY_RUN", "").strip() == "1":
        return

    def worker() -> None:
        try:
            _send_gmail_message(to_address=to_address, subject=subject, body=body)
            try:
                job_pipeline.append_log(
                    "=== VISIT NOTIFY SENT ===\n"
                    f"time: {datetime.now().isoformat(timespec='seconds')}\n"
                    f"page: {page}\n"
                    f"remote: {ip}\n"
                )
            except Exception:
                pass
        except Exception as exc:
            # Don't break page load if email fails.
            try:
                job_pipeline.append_log(
                    "=== VISIT NOTIFY FAILED ===\n"
                    f"time: {datetime.now().isoformat(timespec='seconds')}\n"
                    f"page: {page}\n"
                    f"remote: {ip}\n"
                    f"error: {exc}\n"
                )
            except Exception:
                pass

    threading.Thread(target=worker, daemon=True).start()

# -------------------------
# Serve selected ai-bots web assets
# -------------------------
@app.route("/ai-bots/Thermostats/Web/dashboard_public.html", strict_slashes=False)
def thermostat_dashboard_public():
    return send_from_directory(THERMOSTATS_WEB_DIR, "dashboard_public.html")


@app.route("/ai-bots/Thermostats/Web/<path:filename>")
def thermostats_web_assets(filename):
    return send_from_directory(THERMOSTATS_WEB_DIR, filename)


# -------------------------
# Serve static files (CSS/JS)
# -------------------------
@app.route("/tower/<path:filename>")
def tower_static(filename):
    return send_from_directory(UI_DIR, filename)


@app.route("/tower/read-log")
def read_log():
    try:
        with open(job_pipeline.LOG_FILE_PATH, encoding="utf-8") as f:
            return jsonify({"text": f.read()})
    except FileNotFoundError:
        return jsonify({"text": ""})


@app.route("/tower/clear-log", methods=["POST"])
def clear_log():
    job_pipeline.LOG_FILE_PATH.parent.mkdir(parents=True, exist_ok=True)
    job_pipeline.LOG_FILE_PATH.write_text("", encoding="utf-8")
    return jsonify({"status": "cleared"})


@app.route("/Images/<path:filename>")
def shared_image(filename):
    return send_from_directory(IMAGES_DIR, filename)


@app.route("/view/<path:filename>")
def view_file(filename):
    return send_from_directory(BASE_OUTPUT_DIR, filename, mimetype="application/pdf")


@app.route("/cores/files")
def cores_files():
    try:
        names = sorted(
            [
                path.name
                for path in CORES_DIR.iterdir()
                if path.is_file() and path.suffix.lower() == ".md"
            ],
            key=lambda item: item.lower(),
        )
        return jsonify({"files": names})
    except Exception as exc:
        return jsonify({"error": "Failed to list markdown files", "details": str(exc)}), 500


@app.route("/cores/new-template", methods=["OPTIONS", "POST"])
def cores_new_template():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    try:
        if not CORES_EMPTY_TEMPLATE_PATH.exists():
            return _with_cors(
                app.make_response(
                    (
                        jsonify(
                            {
                                "error": (
                                    f"Template not found: {CORES_EMPTY_TEMPLATE_PATH.name}"
                                )
                            }
                        ),
                        404,
                    )
                )
            )

        template_markdown = CORES_EMPTY_TEMPLATE_PATH.read_text(encoding="utf-8")
        year_value = datetime.now().year

        # Reserve the next available numeric slot for this year.
        attempts = 0
        core_id = ""
        target_path = None
        markdown = ""
        created = False
        while attempts < 1000:
            attempts += 1
            core_id = _next_core_id_for_year(year_value)
            target_path = _resolve_core_markdown_path(core_id)
            markdown = _with_core_id_in_template(template_markdown, core_id)
            try:
                with open(target_path, "x", encoding="utf-8") as fh:
                    fh.write(markdown)
                created = True
                break
            except FileExistsError:
                year_value = datetime.now().year
                continue
        if not created or target_path is None or not core_id:
            raise ValueError("Unable to allocate a new CORE id.")

        _append_cores_log(f"[IMPORT] created empty template {target_path}")
        return _with_cors(
            jsonify(
                {
                    "coreId": core_id,
                    "fileName": target_path.name,
                    "markdown": markdown,
                }
            )
        )
    except PermissionError:
        return _with_cors(
            app.make_response((jsonify({"error": "Path traversal blocked"}), 403))
        )
    except ValueError as exc:
        return _with_cors(
            app.make_response((jsonify({"error": str(exc)}), 422))
        )
    except Exception as exc:
        return _with_cors(
            app.make_response(
                (
                    jsonify({"error": "Failed to create template copy", "details": str(exc)}),
                    500,
                )
            )
        )


@app.route("/CORES", strict_slashes=False)
def cores_index():
    return send_from_directory(CORES_DIR, "CORES.html")


@app.route("/CORES/<core_id>.pdf")
def cores_pdf(core_id):
    md_filename = f"{core_id}.md"
    try:
        base = CORES_DIR.resolve()
        requested_path = (base / md_filename).resolve()
        if not str(requested_path).startswith(str(base)):
            _append_cores_log(
                f"[CORES] pdf path traversal blocked for {core_id} -> {requested_path}"
            )
            abort(403)
        if not requested_path.exists():
            alt_path = (base / f"{core_id}.MD").resolve()
            if alt_path.exists() and str(alt_path).startswith(str(base)):
                requested_path = alt_path
            else:
                _append_cores_log(f"[CORES] pdf missing {requested_path}")
                abort(404)
        _append_cores_log(f"[CORES] pdf generating {requested_path}")
        markdown = requested_path.read_text(encoding="utf-8")
        pdf_bytes = _markdown_to_pdf_bytes(markdown)
        return send_file(
            BytesIO(pdf_bytes),
            mimetype="application/pdf",
            download_name=f"{core_id}.pdf",
        )
    except HTTPException:
        raise
    except Exception as exc:
        _append_cores_log(f"[CORES] pdf generation failed for {core_id}: {exc}")
        return (
            jsonify({"status": "error", "error": "PDF conversion failed"}),
            500,
        )


@app.route("/CORES/<path:filename>")
def cores_static(filename):
    try:
        base = CORES_DIR.resolve()
        requested_path = (base / filename).resolve()
        if not str(requested_path).startswith(str(base)):
            _append_cores_log(
                f"[CORES] path traversal blocked for {filename} -> {requested_path}"
            )
        elif requested_path.exists():
            _append_cores_log(f"[CORES] resolved path {requested_path}")
        else:
            _append_cores_log(f"[CORES] missing file {requested_path}")
    except Exception as exc:
        _append_cores_log(f"[CORES] failed to resolve path for {filename}: {exc}")
    return send_from_directory(CORES_DIR, filename)


@app.route("/import_resume_template", methods=["OPTIONS", "POST"])
def import_resume_template():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    template_markdown = str(request.form.get("templateMarkdown") or "").strip()
    resume_file = request.files.get("resumeFile")
    if not template_markdown:
        _append_cores_log("[IMPORT] rejected: templateMarkdown is required")
        return _with_cors(
            app.make_response((jsonify({"error": "templateMarkdown is required"}), 400))
        )
    if resume_file is None or not str(resume_file.filename or "").strip():
        _append_cores_log("[IMPORT] rejected: resumeFile is required")
        return _with_cors(
            app.make_response((jsonify({"error": "resumeFile is required"}), 400))
        )

    filename = str(resume_file.filename or "").strip()
    suffix = Path(filename).suffix.lower()
    _append_cores_log(f"[IMPORT] started file={filename} type={suffix or 'unknown'}")
    if suffix not in {".pdf", ".docx"}:
        _append_cores_log(f"[IMPORT] rejected: unsupported extension for {filename}")
        return _with_cors(
            app.make_response((jsonify({"error": "Only .pdf and .docx files are supported"}), 400))
        )

    file_bytes = resume_file.read() or b""
    if not file_bytes:
        _append_cores_log(f"[IMPORT] rejected: uploaded file is empty for {filename}")
        return _with_cors(
            app.make_response((jsonify({"error": "Uploaded file is empty"}), 400))
        )

    try:
        if suffix == ".pdf":
            resume_text = _extract_pdf_text(file_bytes)
        else:
            resume_text = _extract_docx_text(file_bytes)
    except ValueError as exc:
        _append_cores_log(f"[IMPORT] resume parsing failed for {filename}: {exc}")
        return _with_cors(
            app.make_response((jsonify({"error": "Resume parsing failed", "details": str(exc)}), 422))
        )

    openai_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not openai_key:
        _append_cores_log("[IMPORT] rejected: OPENAI_API_KEY is missing")
        return _with_cors(
            app.make_response((jsonify({"error": "OPENAI_API_KEY is missing"}), 500))
        )

    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini").strip() or "gpt-4o-mini"
    prompt = (
        f"{RESUME_IMPORT_PROMPT}\n\n"
        f"EMPTY CORES-MD TEMPLATE:\n{template_markdown}\n\n"
        f"RESUME SOURCE TEXT:\n{resume_text}\n"
    )

    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {openai_key}",
            },
            json={
                "model": model,
                "temperature": 0.0,
                "messages": [
                    {"role": "user", "content": prompt},
                ],
                "max_tokens": 3800,
            },
            timeout=90,
        )
    except requests.RequestException as exc:
        _append_cores_log(f"[IMPORT] AI request failed for {filename}: {exc}")
        return _with_cors(
            app.make_response((jsonify({"error": "AI request failed", "details": str(exc)}), 503))
        )

    if response.status_code >= 400:
        _append_cores_log(
            f"[IMPORT] OpenAI request failed for {filename}: status={response.status_code}"
        )
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "OpenAI request failed",
                            "details": response.text,
                        }
                    ),
                    response.status_code,
                )
            )
        )

    data = response.json()
    raw_content = (
        data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
    ).strip()
    if not raw_content:
        _append_cores_log(f"[IMPORT] AI returned empty markdown for {filename}")
        return _with_cors(
            app.make_response((jsonify({"error": "AI returned empty markdown"}), 422))
        )

    output_markdown = raw_content.replace("\r\n", "\n").replace("\r", "\n").strip() + "\n"
    normalized_markdown = _normalize_common_ascii(output_markdown)
    if normalized_markdown != output_markdown:
        _append_cores_log(f"[IMPORT] normalized common non-ASCII typography for {filename}")
    output_markdown = normalized_markdown
    wrapped_markdown = _enforce_resume_import_line_length(output_markdown, max_len=120)
    if wrapped_markdown != output_markdown:
        _append_cores_log(f"[IMPORT] wrapped long lines to <=120 chars for {filename}")
        output_markdown = wrapped_markdown
    try:
        _validate_resume_import_output(template_markdown, output_markdown)
        _validate_cores_markdown(output_markdown)
    except ValueError as exc:
        _append_cores_log(f"[IMPORT] output validation failed for {filename}: {exc}")
        return _with_cors(
            app.make_response(
                (
                    jsonify({"error": "AI output failed resume-template validation", "details": str(exc)}),
                    422,
                )
            )
        )

    _append_cores_log(f"[IMPORT] completed for {filename}")
    return _with_cors(jsonify({"markdown": output_markdown}))


@app.route("/parse_project", methods=["OPTIONS", "POST"])
def parse_project():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    payload = request.get_json(force=True, silent=True) or {}
    project_story = str(payload.get("projectStory") or payload.get("transcript") or "").strip()
    raw_use_chatgpt = payload.get("useChatGPT", False)
    if isinstance(raw_use_chatgpt, str):
        use_chatgpt = raw_use_chatgpt.strip().lower() in {"1", "true", "yes", "on"}
    else:
        use_chatgpt = bool(raw_use_chatgpt)
    raw_prompt_only = payload.get("promptOnly", False)
    if isinstance(raw_prompt_only, str):
        prompt_only = raw_prompt_only.strip().lower() in {"1", "true", "yes", "on"}
    else:
        prompt_only = bool(raw_prompt_only)
    _append_cores_log(
        f"[PARSE_PROJECT] useChatGPT={use_chatgpt} promptOnly={prompt_only} storyChars={len(project_story)}"
    )
    if use_chatgpt and prompt_only:
        return _with_cors(jsonify({"prompt": CORES_FORMATTER_BLOCK}))
    if use_chatgpt:
        if not project_story:
            return _with_cors(
                app.make_response((jsonify({"error": "Project story is required"}), 400))
            )
        prompt = f"""Convert the following PROJECT STORY into a CORES project block.

RULES:
- Do NOT invent information
- Use exact CORES project format
- ASCII only
- No em dashes
- If unknown, leave blank but keep labels

FORMAT:

### PROJECT:

- Situation:
...

- What I did:
- ...

- Tools / systems:
- ...

- Result:
...

- Evidence / artifacts:
- ...

- Notes / caveats:
...

PROJECT STORY:
{project_story}
"""
        return _with_cors(jsonify({"mode": "chatgpt", "prompt": prompt}))

    transcript = project_story
    company_name = str(payload.get("companyName") or "").strip()
    existing_project_titles = payload.get("existingProjectTitles")
    if not isinstance(existing_project_titles, list):
        existing_project_titles = []
    existing_project_titles = [
        str(item or "").strip() for item in existing_project_titles if str(item or "").strip()
    ]
    if not transcript:
        return _with_cors(app.make_response((jsonify({"error": "Transcript is required"}), 400)))

    openai_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not openai_key:
        return _with_cors(
            app.make_response((jsonify({"error": "OPENAI_API_KEY is missing"}), 500))
        )

    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini").strip() or "gpt-4o-mini"
    company_context = f"Target company: {company_name}\n" if company_name else ""
    existing_context = (
        "Existing project titles under this company:\n"
        + "\n".join(f"- {title}" for title in existing_project_titles)
        + "\n"
        if existing_project_titles
        else ""
    )
    prompt = (
        f"{PROJECT_PARSE_PROMPT}\n\n"
        f"{company_context}"
        f"{existing_context}"
        f"Input transcript:\n{transcript}\n"
    )

    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {openai_key}",
            },
            json={
                "model": model,
                "temperature": 0.0,
                "messages": [
                    {"role": "user", "content": prompt},
                ],
                "response_format": {"type": "json_object"},
                "max_tokens": 900,
            },
            timeout=60,
        )
    except requests.RequestException as exc:
        return _with_cors(
            app.make_response((jsonify({"error": "AI request failed", "details": str(exc)}), 503))
        )

    if response.status_code >= 400:
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "OpenAI request failed",
                            "details": response.text,
                        }
                    ),
                    response.status_code,
                )
            )
        )

    data = response.json()
    raw_content = (
        data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
    ).strip()
    if not raw_content:
        return _with_cors(
            app.make_response((jsonify({"error": "AI returned empty project content"}), 422))
        )

    try:
        json_text = _extract_first_json_object(raw_content)
        project_json = json.loads(json_text)
        normalized_project = _normalize_project_json(project_json)
        project_block = _build_project_block(
            normalized_project["projectTitle"],
            normalized_project["sections"],
        )
        _parse_project_block_sections(project_block)
    except json.JSONDecodeError as exc:
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "AI output failed CORES project validation",
                            "details": f"AI JSON decode failed: {exc}",
                        }
                    ),
                    422,
                )
            )
        )

    except ValueError as exc:
        return _with_cors(
            app.make_response(
                (
                    jsonify({"error": "AI output failed CORES project validation", "details": str(exc)}),
                    422,
                )
            )
        )

    return _with_cors(jsonify({"project": normalized_project, "projectBlock": project_block}))


@app.route("/parse_company", methods=["OPTIONS", "POST"])
def parse_company():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    payload = request.get_json(force=True, silent=True) or {}
    raw_use_chatgpt = payload.get("useChatGPT", False)
    if isinstance(raw_use_chatgpt, str):
        use_chatgpt = raw_use_chatgpt.strip().lower() in {"1", "true", "yes", "on"}
    else:
        use_chatgpt = bool(raw_use_chatgpt)
    raw_prompt_only = payload.get("promptOnly", False)
    if isinstance(raw_prompt_only, str):
        prompt_only = raw_prompt_only.strip().lower() in {"1", "true", "yes", "on"}
    else:
        prompt_only = bool(raw_prompt_only)
    _append_cores_log(
        f"[PARSE_COMPANY] useChatGPT={use_chatgpt} promptOnly={prompt_only}"
    )
    transcript = str(payload.get("transcript") or "").strip()
    existing_company_names = payload.get("existingCompanyNames")
    if not isinstance(existing_company_names, list):
        existing_company_names = []
    existing_company_names = [
        str(item or "").strip() for item in existing_company_names if str(item or "").strip()
    ]
    if use_chatgpt and prompt_only:
        try:
            company_template_markdown = _load_company_template_markdown()
            formatter_block = _build_company_formatter_block(company_template_markdown)
            return _with_cors(jsonify({"prompt": formatter_block}))
        except ValueError as exc:
            _append_cores_log(f"[PARSE_COMPANY] template load failed: {exc}")
            return _with_cors(app.make_response((jsonify({"error": str(exc)}), 422)))
        except Exception as exc:
            _append_cores_log(f"[PARSE_COMPANY] template read failed: {exc}")
            return _with_cors(
                app.make_response((jsonify({"error": "Failed to read company template", "details": str(exc)}), 500))
            )
    if not transcript:
        return _with_cors(app.make_response((jsonify({"error": "Transcript is required"}), 400)))

    openai_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not openai_key:
        return _with_cors(
            app.make_response((jsonify({"error": "OPENAI_API_KEY is missing"}), 500))
        )

    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini").strip() or "gpt-4o-mini"
    try:
        company_template_markdown = _load_company_template_markdown()
    except ValueError as exc:
        _append_cores_log(f"[PARSE_COMPANY] template load failed: {exc}")
        return _with_cors(app.make_response((jsonify({"error": str(exc)}), 422)))
    except Exception as exc:
        _append_cores_log(f"[PARSE_COMPANY] template read failed: {exc}")
        return _with_cors(
            app.make_response((jsonify({"error": "Failed to read company template", "details": str(exc)}), 500))
        )
    existing_context = (
        "Existing company names:\n"
        + "\n".join(f"- {name}" for name in existing_company_names)
        + "\n"
        if existing_company_names
        else ""
    )
    prompt = (
        f"{COMPANY_PARSE_PROMPT}\n\n"
        "LOCKED COMPANY FRAMEWORK (reference for mapping fields):\n"
        f"{company_template_markdown}\n\n"
        "Use the story to populate this framework mapping in JSON only:\n"
        '- "companyName" = value from "## COMPANY: ..."\n'
        '- "role" = content under "### ROLE:"\n'
        '- "signals" = content under "### SIGNALS:"\n'
        '- "firstProject.projectTitle" = optional text after "### PROJECT:"\n'
        '- "firstProject.sections" = Situation, What I did, Tools / systems, Result, Evidence / artifacts, Notes / caveats\n\n'
        f"{existing_context}"
        f"Input transcript:\n{transcript}\n"
    )

    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {openai_key}",
            },
            json={
                "model": model,
                "temperature": 0.0,
                "messages": [
                    {"role": "user", "content": prompt},
                ],
                "response_format": {"type": "json_object"},
                "max_tokens": 1300,
            },
            timeout=60,
        )
    except requests.RequestException as exc:
        return _with_cors(
            app.make_response((jsonify({"error": "AI request failed", "details": str(exc)}), 503))
        )

    if response.status_code >= 400:
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "OpenAI request failed",
                            "details": response.text,
                        }
                    ),
                    response.status_code,
                )
            )
        )

    data = response.json()
    raw_content = (
        data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
    ).strip()
    if not raw_content:
        return _with_cors(
            app.make_response((jsonify({"error": "AI returned empty company content"}), 422))
        )

    try:
        json_text = _extract_first_json_object(raw_content)
        company_json = json.loads(json_text)
        normalized_company = _normalize_company_json(company_json)
        project_block = _build_project_block(
            normalized_company["firstProject"]["projectTitle"],
            normalized_company["firstProject"]["sections"],
        )
        _parse_project_block_sections(project_block)
    except json.JSONDecodeError as exc:
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "AI output failed CORES company validation",
                            "details": f"AI JSON decode failed: {exc}",
                        }
                    ),
                    422,
                )
            )
        )

    except ValueError as exc:
        return _with_cors(
            app.make_response(
                (
                    jsonify({"error": "AI output failed CORES company validation", "details": str(exc)}),
                    422,
                )
            )
        )

    return _with_cors(jsonify({"company": normalized_company, "projectBlock": project_block}))


@app.route("/save_cores", methods=["OPTIONS", "POST"])
def save_cores():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    payload = request.get_json(force=True, silent=True) or {}
    core_id = str(payload.get("coreId") or "").strip()
    markdown = payload.get("markdown")
    if not core_id:
        return _with_cors(app.make_response((jsonify({"error": "coreId is required"}), 400)))
    if not isinstance(markdown, str):
        return _with_cors(app.make_response((jsonify({"error": "markdown text is required"}), 400)))

    try:
        _validate_cores_markdown(markdown)
        target_path = _resolve_core_markdown_path(core_id)
    except PermissionError:
        return _with_cors(app.make_response((jsonify({"error": "Path traversal blocked"}), 403)))
    except ValueError as exc:
        return _with_cors(app.make_response((jsonify({"error": str(exc)}), 422)))

    if not target_path.exists():
        return _with_cors(
            app.make_response((jsonify({"error": f"Source file not found: {target_path.name}"}), 404))
        )

    try:
        target_path.write_text(markdown, encoding="utf-8")
        _append_cores_log(f"[CORES] saved {target_path}")
    except Exception as exc:
        return _with_cors(
            app.make_response((jsonify({"error": "Save failed", "details": str(exc)}), 500))
        )

    return _with_cors(
        jsonify({"status": "ok", "saved": target_path.name})
    )


# -------------------------
# BUTTON ENDPOINTS (STUB ONLY)
# Each returns a simple JSON message
# -------------------------

@app.route("/tower/import-job")
def import_job():
    job_pipeline.run_import_job()
    return jsonify({"status": "launched"})


@app.route("/tower/scam-check")
def scam_check():
    job_pipeline.run_scam_check()
    return jsonify({"status": "launched"})


@app.route("/tower/verify-company")
def verify_company():
    job_pipeline.run_verify_company()
    return jsonify({"status": "launched"})


@app.route("/tower/jason-configuration")
def jason_configuration():
    job_pipeline.run_jason_configuration()
    return jsonify({"status": "launched"})


@app.route("/tower/facility-check")
def facility_check():
    job_pipeline.run_facility_check()
    return jsonify({"status": "launched"})


@app.route("/tower/enrich-company")
def enrich_company():
    job_pipeline.stub()
    return jsonify({"status": "ok", "action": "enrich_company stub called"})


@app.route("/tower/targeted-resume")
def targeted_resume():
    job_pipeline.run_targeted_resume()
    return jsonify({"status": "launched"})


@app.route("/tower/targeted-resume/variant/<int:variant>")
def targeted_resume_variant_route(variant):
    try:
        job_pipeline.run_targeted_resume_variant(variant)
    except ValueError as exc:
        return jsonify({"status": "error", "error": str(exc)}), 400
    return jsonify({"status": "launched", "variant": variant})


@app.route("/tower/yahoo-forwarder")
def yahoo_forwarder():
    job_pipeline.run_yahoo_forwarder()
    return jsonify({"status": "launched"})


@app.route("/tower/craigslist-forwarder")
def craigslist_forwarder():
    # Backward-compatible alias for the renamed YahooForwarder project.
    job_pipeline.run_yahoo_forwarder()
    return jsonify({"status": "launched"})


@app.route("/tower/push-ss")
def push_ss():
    job_pipeline.stub()
    return jsonify({"status": "ok", "action": "push_ss stub called"})


@app.route("/tower/read-thermostat")
def read_thermostat():
    thermostat.stub()
    return jsonify({"status": "ok", "action": "read_thermostat stub called"})


@app.route("/tower/open-dashboard")
def open_dashboard():
    thermostat.stub()
    return jsonify({"status": "ok", "action": "open_dashboard stub called"})


@app.route("/tower/weather-sync")
def weather_sync():
    thermostat.stub()
    return jsonify({"status": "ok", "action": "weather_sync stub called"})


@app.route("/tower/fax-queue")
def fax_queue():
    fax_engine.stub()
    return jsonify({"status": "ok", "action": "fax_queue stub called"})


@app.route("/tower/fax-retry")
def fax_retry():
    fax_engine.stub()
    return jsonify({"status": "ok", "action": "fax_retry stub called"})


@app.route("/tower/fax-log")
def fax_log():
    fax_engine.stub()
    return jsonify({"status": "ok", "action": "fax_log stub called"})


@app.route("/tower/security-notify")
def security_notify_route():
    security_notify.stub()
    return jsonify({"status": "ok", "action": "security_notify stub called"})


@app.route("/tower/security-log")
def security_log():
    security_notify.stub()
    return jsonify({"status": "ok", "action": "security_log stub called"})


@app.route("/tower/config-load")
def config_load():
    config_manager.stub()
    return jsonify({"status": "ok", "action": "config_load stub called"})


@app.route("/tower/config-edit")
def config_edit():
    config_manager.stub()
    return jsonify({"status": "ok", "action": "config_edit stub called"})


@app.route("/tower/config-refresh")
def config_refresh():
    config_manager.stub()
    return jsonify({"status": "ok", "action": "config_refresh stub called"})


@app.route("/tower/stinky-bath")
def stinky_bath():
    logger.stub()
    return jsonify({"status": "ok", "action": "stinky_bath stub called"})


@app.route("/tower/scheduled-tasks")
def scheduled_tasks():
    try:
        return jsonify({"status": "ok", "tasks": task_scheduler.list_hsst_tasks_payload()})
    except Exception as exc:
        return jsonify({"status": "error", "error": str(exc), "tasks": []}), 500


def _build_prompt(
    question: str,
    history: str,
    job_description: str = "",
    skills: str = "",
    supplemental_context: str = "",
    mode: str = "summary",
) -> str:
    job_description = (job_description or "").strip()
    history = (history or "").strip()
    skills = (skills or "").strip()
    supplemental_context = (supplemental_context or "").strip()
    question = (question or "").strip()
    mode = (mode or "summary").strip().lower()

    skills_block = f"\n\nSelf-rated skills:\n{skills}\n" if skills else ""
    supplemental_block = (
        f"\n\nSupplemental context (candidate-provided):\n{supplemental_context}\n"
        if supplemental_context
        else ""
    )

    detail_instruction = ""
    if mode == "deep":
        detail_instruction = (
            "\n\nDepth mode: DEEP\n"
            "- Provide a thorough, multi-angle analysis.\n"
            "- Use more bullets, more specifics, and more evidence.\n"
            "- Prefer concrete examples (projects, metrics, tools, workflows).\n"
        )

    if job_description:
        return (
            "You are a candid, skeptical career evaluator. Do not be flattering.\n"
            "Rules:\n"
            "- Only claim skills/experience that are supported by the work history OR explicitly listed in the self-rated skills OR described in the supplemental context.\n"
            "- If a skill is only in the self-rated skills (not in work history), label it as self-rated.\n"
            "- If a claim is only supported by the supplemental context (not work history), label it as candidate-provided.\n"
            "- If the job is unrelated to the work history, say so plainly.\n"
            "- Prefer \"No\" or \"Maybe\" over a false \"Yes\".\n"
            "- Tie every positive claim to explicit evidence from the work history when possible.\n\n"
            "Important:\n"
            "- Before writing the answer, scan the supplemental context for any lines that look like project headers (e.g., starting with \"Project:\").\n"
            "- Use those project names verbatim when relevant (example: \"TileAir / DocOrigin Replacement Initiative\").\n\n"
            "Task:\n"
            "Evaluate whether the candidate is a good fit for the job description below.\n"
            "Return Markdown with this exact structure:\n"
            "## Fit Verdict\n"
            "- Verdict: Yes/Maybe/No\n"
            "- Confidence: 0-100\n"
            "## Evidence From Work History\n"
            "- (3-8 bullets) Each bullet must reference something from the work history.\n"
            "## Relevant Projects\n"
            "- (5-12 bullets) Include specific projects/initiatives; prefer exact names from the supplemental context when available.\n"
            "- Each bullet must end with a support tag: (work history) or (candidate-provided) or (self-rated).\n"
            "## Gaps / Risks\n"
            "- (3-8 bullets)\n"
            "## Interview Positioning\n"
            "- (3-8 bullets)\n\n"
            f"{detail_instruction}\n"
            "Job description:\n"
            f"{job_description}\n\n"
            "Work history:\n"
            f"{history}{skills_block}{supplemental_block}\n"
            f"Question: {question}"
        )

    return (
        "You are a helpful but honest career advisor.\n"
        "Rules:\n"
        "- Only claim skills/experience that are supported by the work history OR explicitly listed in the self-rated skills OR described in the supplemental context.\n"
        "- If a skill is only in the self-rated skills (not in work history), label it as self-rated.\n"
        "- If a claim is only supported by the supplemental context (not work history), label it as candidate-provided.\n"
        "- If you don't know, say so.\n\n"
        "Use the work history, self-rated skills, and supplemental context below to answer the user's question. "
        "Focus on automation, reliability, and human-led systems.\n"
        "If the user explicitly requests a single sentence, comply.\n"
        "Otherwise, provide at least three supporting points.\n\n"
        f"{detail_instruction}\n"
        "Work history:\n"
        f"{history}{skills_block}{supplemental_block}\n"
        f"Question: {question}"
    )


@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET,POST,OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return response


@app.route("/tower/contact", methods=["OPTIONS", "POST"])
def tower_contact():
    if request.method == "OPTIONS":
        return ("", 204)

    payload = request.get_json(force=True, silent=True) or {}
    website = (payload.get("website") or "").strip()
    name = (payload.get("name") or "").strip()
    contact = (payload.get("contact") or "").strip()
    message = (payload.get("message") or "").strip()
    page_url = (payload.get("pageUrl") or "").strip()

    # Honeypot for bots (humans won't see/fill this).
    if website:
        return jsonify({"status": "ok"})

    ip = (request.remote_addr or "").strip() or "unknown"
    now = time.time()
    last = _contact_last_sent_at.get(ip)
    if last is not None and now - last < 30:
        return jsonify({"error": "Too many requests. Try again shortly."}), 429

    if not contact:
        return jsonify({"error": "Contact info is required"}), 400

    if len(name) > 120:
        name = name[:120]
    if len(contact) > 220:
        contact = contact[:220]
    if len(message) > 5000:
        message = message[:5000]
    if len(page_url) > 1000:
        page_url = page_url[:1000]

    to_address = os.environ.get("HSST_CONTACT_TO", "timinman2024@gmail.com").strip()
    reply_to = contact if _looks_like_email(contact) else None

    subject_hint = name or contact
    subject = f"AI-Fit-Site contact: {subject_hint}".strip()

    body_lines = [
        "New contact request from AI-Fit-Site",
        "",
        f"Time: {datetime.now().isoformat(timespec='seconds')}",
        f"From (provided): {name or '(no name)'}",
        f"Contact (provided): {contact}",
        f"Page: {page_url}" if page_url else "",
        "",
        "Message:",
        message or "(no message)",
        "",
        f"Remote: {request.remote_addr}",
        f"Origin: {request.headers.get('Origin', '').strip()}",
        f"User-Agent: {request.headers.get('User-Agent', '').strip()}",
    ]
    body = "\n".join(body_lines).strip() + "\n"

    if os.environ.get("HSST_CONTACT_DRY_RUN", "").strip() == "1":
        _contact_last_sent_at[ip] = now
        return jsonify({"status": "ok", "dryRun": True})

    try:
        _send_gmail_message(to_address=to_address, subject=subject, body=body, reply_to=reply_to)
    except Exception as exc:
        return jsonify({"error": "Failed to send email", "details": str(exc)}), 500

    _contact_last_sent_at[ip] = now
    return jsonify({"status": "ok"})


@app.route("/tower/ask-ai", methods=["OPTIONS", "POST"])
def ask_ai():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    payload = request.get_json(force=True, silent=True) or {}
    question = (payload.get("question") or "").strip()
    mode = (payload.get("mode") or "summary").strip().lower()
    history = payload.get("history", "").strip()
    skills = payload.get("skills", "").strip()
    supplemental_context = payload.get("supplementalContext", "").strip()
    job_description = (payload.get("jobDescription") or "").strip()

    if mode not in {"summary", "deep"}:
        mode = "summary"

    job_pipeline.append_log(
        "=== Ask AI REQUEST ===\n"
        f"time: {datetime.now().isoformat(timespec='seconds')}\n"
        f"remote: {request.remote_addr}\n"
        f"origin: {request.headers.get('Origin', '')}\n"
        f"ua: {request.headers.get('User-Agent', '')}\n"
        f"mode: {mode}\n"
        f"question_preview: {question[:160].replace(chr(10), ' ')!r}\n"
        f"history_chars: {len(history)}\n"
        f"skills_chars: {len(skills)}\n"
        f"supplemental_chars: {len(supplemental_context)}\n"
        f"job_desc_chars: {len(job_description)}\n"
    )
    if not question:
        job_pipeline.append_log("=== Ask AI ERROR ===\nerror: missing question\n")
        return _with_cors(app.make_response((jsonify({"error": "Question is required"}), 400)))

    openai_key = os.environ.get("OPENAI_API_KEY")
    if not openai_key:
        job_pipeline.append_log("=== Ask AI ERROR ===\nerror: OPENAI_API_KEY missing\n")
        return _with_cors(
            app.make_response(
                (
                    jsonify(
                        {
                            "error": "OPENAI_API_KEY is missing. Set it in the environment before running tower.",
                        }
                    ),
                    500,
                )
            )
        )

    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
    prompt = _build_prompt(question, history, job_description, skills, supplemental_context, mode)
    max_tokens = 950 if mode == "summary" else 1600
    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {openai_key}",
            },
            json={
                "model": model,
                "temperature": 0.1,
                "messages": [
                    {
                        "role": "system",
                        "content": (
                            "You are a career advisor who must be honest, specific, and evidence-based. "
                            "Do not overstate fit or invent experience."
                        ),
                    },
                    {
                        "role": "user",
                        "content": prompt,
                    },
                ],
                "max_tokens": max_tokens,
            },
            timeout=60,
        )
        if response.status_code >= 400:
            job_pipeline.append_log(
                "=== Ask AI ERROR ===\n"
                f"error: OpenAI request failed ({response.status_code})\n"
                f"details_preview: {response.text[:600]!r}\n"
            )
            return _with_cors(
                app.make_response(
                    (
                        jsonify(
                            {
                                "error": "OpenAI request failed",
                                "details": response.text,
                            }
                        ),
                        response.status_code,
                    )
                )
            )

        data = response.json()
        answer = data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()
        job_pipeline.append_log(
            "=== Ask AI RESPONSE ===\n"
            f"answer_chars: {len(answer)}\n"
        )
        return _with_cors(app.make_response(jsonify({"answer": answer or "OpenAI returned an empty response."})))
    except requests.RequestException as exc:
        job_pipeline.append_log(
            "=== Ask AI ERROR ===\n"
            f"error: Proxy request failed\n"
            f"details: {str(exc)!r}\n"
        )
        return _with_cors(
            app.make_response((jsonify({"error": "Proxy request failed", "details": str(exc)}), 503))
        )


@app.route("/cores/log", methods=["OPTIONS", "POST"])
def cores_log():
    if request.method == "OPTIONS":
        resp = app.make_response(("", 204))
        return _with_cors(resp)

    payload = request.get_json(force=True, silent=True) or {}
    message = (payload.get("message") or "").strip()
    if not message:
        return _with_cors(
            app.make_response(
                (
                    jsonify({"error": "Message is required"}),
                    400,
                )
            )
        )

    _append_cores_log(message)
    return _with_cors(jsonify({"status": "ok"}))

# -------------------------
# RUN SERVER
# -------------------------
def _seconds_until_next_snapshot() -> float:
    now = datetime.now()
    next_run = now.replace(hour=0, minute=1, second=0, microsecond=0)
    if next_run <= now:
        next_run += timedelta(days=1)
    return (next_run - now).total_seconds()


def run_dashboard_snapshot() -> None:
    label = "Nightly Dashboard Snapshot"
    job_pipeline.append_log(f"=== {label} STARTED ===\n")
    try:
        result = subprocess.run(
            [sys.executable, str(DASHBOARD_SCRIPT_PATH)],
            capture_output=True,
            text=True,
        )
        if result.stdout:
            job_pipeline.append_log(result.stdout)
        if result.stderr:
            job_pipeline.append_log(f"[stderr] {result.stderr}")
        if result.returncode != 0:
            raise RuntimeError(
                f"{DASHBOARD_SCRIPT_PATH} exited with {result.returncode}"
            )
    except Exception as exc:
        job_pipeline.append_log(f"=== {label} FAILED ===\n{exc}\n")
    finally:
        job_pipeline.append_log(f"=== {label} FINISHED ===\n")


def schedule_dashboard_snapshot() -> None:
    def worker() -> None:
        while True:
            wait_seconds = _seconds_until_next_snapshot()
            time.sleep(wait_seconds)
            run_dashboard_snapshot()

    threading.Thread(target=worker, daemon=True).start()


if __name__ == "__main__":
    schedule_dashboard_snapshot()
    print("HSST Control Tower Flask Server Running (Skeleton Mode)")
    app.run(host="0.0.0.0", port=5000, debug=True)
